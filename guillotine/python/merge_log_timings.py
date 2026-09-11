"""
merge_log_timings.py

Extracts generation timings and timeout events from the batch runner log
and merges them into the results workbook:
- adds duration_s, is_first_call, status columns to existing rows
- inserts explicit rows for (model, T, prompt) combinations that timed
  out, so far absent from the sheet as a silent gap (no row at all)

is_first_call=1 identifies, for each model, the chronologically earliest
successful generation in the log - the one whose duration_s (likely)
includes loading the model into VRAM, not just inference. It should be
excluded or replaced with an average before using duration_s as a proxy
for "how much the model had to reason".

Usage:
    python merge_log_timings.py D:\\guillotine\\guillotine\\source\\batch_runner.log D:\\guillotine\\guillotine\\source\\ghigliottina_results.xlsx
"""
import argparse
import json
import logging
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

import config

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

OK_PATTERN = re.compile(
    r"^(?P<ts>[\d-]+ [\d:,]+) - INFO - \[ok\] (?P<model>\S+) \| T=(?P<temp>[\d.]+) \| "
    r"(?P<prompt_id>\S+) -> .*?, (?P<duration>[\d.]+)s\)\s*$"
)
FAIL_PATTERN = re.compile(
    r"^(?P<ts>[\d-]+ [\d:,]+) - ERROR - \[fail\] (?P<model>\S+) \| T=(?P<temp>[\d.]+) \| "
    r"(?P<prompt_id>\S+): (?P<error>.+)$"
)
SHEET_BY_TEMP_STR = {"0.0": "T_0.0", "0.3": "T_0.3", "0.7": "T_0.7"}
NEW_COLUMNS = ["duration_s", "is_first_call", "status"]


@dataclass
class LogEntry:
    status: str  # "ok" | "timeout" | "error"
    duration_s: float
    timestamp: str


def parse_log(log_path: Path) -> dict[tuple[str, str, str], LogEntry]:
    """Returns {(model, sheet_name, prompt_id): LogEntry}.

    If the same combination has multiple attempts in the log (e.g. a
    timeout followed by a successful retry after a restart), the
    chronologically latest occurrence wins - reflecting the final state,
    consistent with what actually ends up (or doesn't) in the Excel file.
    """
    entries: dict[tuple[str, str, str], LogEntry] = {}

    with open(log_path, "r", encoding="utf-8", errors="replace") as f:
        for raw_line in f:
            line = raw_line.rstrip("\r\n")

            m = OK_PATTERN.match(line)
            if m:
                sheet = SHEET_BY_TEMP_STR.get(m.group("temp"))
                if sheet is None:
                    logger.warning(f"Unrecognized temperature: {line[:100]!r}")
                    continue
                key = (m.group("model"), sheet, m.group("prompt_id"))
                entries[key] = LogEntry(
                    status="ok", duration_s=float(m.group("duration")), timestamp=m.group("ts")
                )
                continue

            m = FAIL_PATTERN.match(line)
            if m:
                sheet = SHEET_BY_TEMP_STR.get(m.group("temp"))
                if sheet is None:
                    logger.warning(f"Unrecognized temperature: {line[:100]!r}")
                    continue
                key = (m.group("model"), sheet, m.group("prompt_id"))
                status = "timeout" if "timeout" in m.group("error").lower() else "error"
                entries[key] = LogEntry(
                    status=status,
                    duration_s=float(config.OLLAMA_GENERATE_TIMEOUT_S),
                    timestamp=m.group("ts"),
                )

    n_ok = sum(1 for e in entries.values() if e.status == "ok")
    n_timeout = sum(1 for e in entries.values() if e.status == "timeout")
    n_error = sum(1 for e in entries.values() if e.status == "error")
    logger.info(f"Log parsed: {n_ok} ok, {n_timeout} timeout, {n_error} other errors")
    return entries


def first_call_keys(entries: dict[tuple[str, str, str], LogEntry]) -> set[tuple[str, str, str]]:
    """For each model, the chronologically earliest 'ok' entry in the log."""
    earliest: dict[str, tuple[str, tuple[str, str, str]]] = {}
    for key, entry in entries.items():
        if entry.status != "ok":
            continue
        model = key[0]
        if model not in earliest or entry.timestamp < earliest[model][0]:
            earliest[model] = (entry.timestamp, key)
    return {v[1] for v in earliest.values()}


def load_prompts_by_id(prompts_json_path: Path) -> dict[str, dict]:
    with open(prompts_json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    return {p["prompt_id"]: p for p in data["prompts"]}


def ensure_columns(ws) -> dict[str, int]:
    """Appends the new columns at the end if not already present, returns the updated column index."""
    headers = [c.value for c in ws[1]]
    for col_name in NEW_COLUMNS:
        if col_name not in headers:
            ws.cell(row=1, column=len(headers) + 1, value=col_name)
            headers.append(col_name)
    return {name: idx + 1 for idx, name in enumerate(headers)}


def merge_into_workbook(
    xlsx_path: Path, entries: dict[tuple[str, str, str], LogEntry], prompts_by_id: dict[str, dict]
) -> None:
    wb = load_workbook(xlsx_path)
    first_calls = first_call_keys(entries)

    matched = 0
    unmatched_existing = 0
    inserted_timeouts = 0

    for sheet_name in ["T_0.0", "T_0.3", "T_0.7"]:
        ws = wb[sheet_name]
        col_index = ensure_columns(ws)
        model_col = col_index["model_name"]
        prompt_col = col_index["prompt_id"]

        present_pairs: set[tuple[str, str]] = set()

        for row_idx in range(2, ws.max_row + 1):
            model = ws.cell(row=row_idx, column=model_col).value
            prompt_id = ws.cell(row=row_idx, column=prompt_col).value
            if not model or not prompt_id:
                continue
            present_pairs.add((model, prompt_id))

            entry = entries.get((model, sheet_name, prompt_id))
            if entry is None:
                unmatched_existing += 1
                continue

            ws.cell(row=row_idx, column=col_index["duration_s"], value=entry.duration_s)
            ws.cell(
                row=row_idx,
                column=col_index["is_first_call"],
                value=int((model, sheet_name, prompt_id) in first_calls),
            )
            ws.cell(row=row_idx, column=col_index["status"], value=entry.status)
            matched += 1

        # Explicit rows for the timeouts/errors: present in the log but
        # absent from the sheet. Until now a silent gap (no row) - here
        # they become an actual data point.
        temp_str = sheet_name.replace("T_", "")
        for (model, sheet, prompt_id), entry in entries.items():
            if sheet != sheet_name or entry.status == "ok":
                continue
            if (model, prompt_id) in present_pairs:
                continue

            prompt_info = prompts_by_id.get(prompt_id)
            if prompt_info is None:
                logger.warning(f"Unknown prompt_id, timeout row not inserted: {prompt_id}")
                continue

            next_row = ws.max_row + 1
            values = {
                "model_name": model,
                "prompt_id": prompt_id,
                "word1": prompt_info["clues"][0],
                "word2": prompt_info["clues"][1],
                "word3": prompt_info["clues"][2],
                "word4": prompt_info["clues"][3],
                "word5": prompt_info["clues"][4],
                "expected_solution": prompt_info["solution"],
                "raw_response": f"[{entry.status.upper()}]",
                "predicted_word": "",
                "format_compliant": 0,
                "exact_match": 0,
                "same_root": 0,
                "levenshtein_normalized": 1.0,
                "is_similar_lev": 0,
                "timestamp": entry.timestamp,
                "difficulty_note": prompt_info["difficulty"],
                "duration_s": entry.duration_s,
                "is_first_call": 0,
                "status": entry.status,
            }
            for col_name, value in values.items():
                if col_name in col_index:
                    ws.cell(row=next_row, column=col_index[col_name], value=value)
            inserted_timeouts += 1
            present_pairs.add((model, prompt_id))

    wb.save(xlsx_path)
    logger.info(
        f"Existing rows updated with duration/status: {matched} | "
        f"present in the sheet but with no match in the log: {unmatched_existing} | "
        f"timeout/error rows inserted: {inserted_timeouts}"
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Merges timings and timeouts from the log into the results workbook"
    )
    parser.add_argument("log_path", type=Path)
    parser.add_argument("xlsx_path", type=Path)
    parser.add_argument(
        "--prompts-json", type=Path, default=None,
        help="Default: config.PROMPTS_JSON",
    )
    args = parser.parse_args()

    if not args.log_path.exists():
        logger.error(f"Log not found: {args.log_path}")
        sys.exit(1)
    if not args.xlsx_path.exists():
        logger.error(f"Workbook not found: {args.xlsx_path}")
        sys.exit(1)

    prompts_json_path = args.prompts_json or config.PROMPTS_JSON
    prompts_by_id = load_prompts_by_id(prompts_json_path)

    entries = parse_log(args.log_path)
    merge_into_workbook(args.xlsx_path, entries, prompts_by_id)


if __name__ == "__main__":
    main()
