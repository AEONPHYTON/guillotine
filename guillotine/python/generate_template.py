"""
generate_template.py

Creates the ghigliottina_results.xlsx workbook with the 4 agreed sheets:
- T_0.0, T_0.3, T_0.7: header only, the batch runner writes the rows for
  the local models.
- Commercial: pre-filled with model_name x prompt_id x clue x expected_solution.
  The raw_response column (highlighted in yellow) must be filled in by hand.

Usage:
    python source/generate_template.py source/guillotine_prompts.json source/ghigliottina_results.xlsx
"""
import sys
import json
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import config

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

RUN_SHEETS = ["T_0.0", "T_0.3", "T_0.7"]

# Commercial models from the study project. Free-form list: add/remove
# rows directly in the Commercial sheet after generation, no need to
# rerun the script for that.
COMMERCIAL_MODELS = [
    "Gemini 3 veloce",
    "Perplexity",
    "Claude Sonnet 4.5",
    "Claude Sonnet 4.6",
    "ChatGPT 5.3",
    "DeepSeek Instant",
    "Qwen 3.6 Plus automatico",
    "Qwen 3.5 Plus automatico",
]

# Single schema lives in config.py - don't duplicate it here. It used to be
# duplicated and silently went out of sync when duration_s/is_first_call/
# status were added to the schema (the same double-maintenance issue
# discussed for Excel/JSON).
COLUMNS = config.RESULT_COLUMNS

# Columns the user has to fill in by hand - highlighted in yellow in the header
FILL_IN_COLUMNS = {"raw_response"}

HEADER_FONT = Font(name="Arial", bold=True)
FILL_IN_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


def write_header(ws) -> None:
    for col_idx, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = HEADER_FONT
        if name in FILL_IN_COLUMNS:
            cell.fill = FILL_IN_FILL
    ws.freeze_panes = "A2"


def autosize(ws) -> None:
    for col_idx, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(name) + 2)


def build_run_sheet(ws) -> None:
    """T_0.0 / T_0.3 / T_0.7: header only, rows added by the batch runner."""
    write_header(ws)
    autosize(ws)


def build_commercial_sheet(ws, prompts: list[dict]) -> None:
    """Pre-fills model x prompt, leaving raw_response and the scores blank."""
    write_header(ws)

    row_idx = 2
    for model_name in COMMERCIAL_MODELS:
        for prompt in prompts:
            values = {
                "model_name": model_name,
                "prompt_id": prompt["prompt_id"],
                "word1": prompt["clues"][0],
                "word2": prompt["clues"][1],
                "word3": prompt["clues"][2],
                "word4": prompt["clues"][3],
                "word5": prompt["clues"][4],
                "expected_solution": prompt["solution"],
                "difficulty_note": prompt["difficulty"],
            }
            for col_idx, name in enumerate(COLUMNS, start=1):
                if name in values:
                    ws.cell(row=row_idx, column=col_idx, value=values[name])
            row_idx += 1

    autosize(ws)
    total_rows = row_idx - 2
    logger.info(
        f"Commercial: {total_rows} rows pre-filled "
        f"({len(COMMERCIAL_MODELS)} models x {len(prompts)} prompts)"
    )


def main() -> None:
    if len(sys.argv) != 3:
        print("Usage: python generate_template.py <guillotine_prompts.json> <output.xlsx>")
        sys.exit(1)

    json_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])

    try:
        with open(json_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except FileNotFoundError:
        logger.error(f"File not found: {json_path}")
        sys.exit(1)
    except json.JSONDecodeError as e:
        logger.error(f"Invalid JSON in {json_path}: {e}")
        sys.exit(1)

    prompts = data["prompts"]
    if not prompts:
        logger.error("No prompts found in the JSON")
        sys.exit(1)

    wb = Workbook()
    wb.remove(wb.active)  # removes the default "Sheet"

    for sheet_name in RUN_SHEETS:
        ws = wb.create_sheet(sheet_name)
        build_run_sheet(ws)

    ws_commercial = wb.create_sheet("Commercial")
    build_commercial_sheet(ws_commercial, prompts)

    wb.save(output_path)
    logger.info(f"Template saved to {output_path}")


if __name__ == "__main__":
    main()
